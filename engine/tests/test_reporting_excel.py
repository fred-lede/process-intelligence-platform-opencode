"""Guard: every populated ReportData section must reach the Excel workbook.

The generator used to write 7 sheets while the HTML report rendered ~20
sections, so data `main.py` had already collected (specifications, SPC
capability, coefficients, credibility, ...) was silently dropped from the
workbook. A missing section is invisible -- the file opens fine -- so this is
asserted mechanically rather than by eye.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from process_intelligence_engine.reporting.excel import ExcelReportGenerator
from process_intelligence_engine.reporting.models import ReportData


def sheet_names(data: ReportData) -> list[str]:
    wb = load_workbook(BytesIO(ExcelReportGenerator(data).generate()))
    return wb.sheetnames


def populated_report() -> ReportData:
    """ReportData with every renderable section filled in."""
    return ReportData(
        project_name="Probe",
        dataset_id="ds-1",
        source_file="probe.csv",
        row_count=90,
        column_count=3,
        time_range={"start": "2026-01-01", "end": "2026-01-31"},
        fields=[{"name": "y", "role": "output", "confidence": 0.9, "data_type": "float"}],
        spec={"limits": {"lsl": 1.56, "usl": 1.68, "target": 1.62}, "y": "mm"},
        quality_summary={
            "issue_count": 2,
            "issues_by_severity": {"warning": 2},
            "issues": [{"check": "missing", "column": "y", "severity": "warning", "message": "1 missing"}],
        },
        distribution_fits={"y": [{"name": "norm", "aic": 1.0, "ks_p_value": 0.4,
                                  "skewness": 0.1, "kurtosis": -0.2}]},
        anomalies=[{"name": "high temp", "type": "engineering", "target_input": "t",
                    "direction": "above", "threshold": 200.0,
                    "occurrence_probability": 0.05, "confidence": 0.8}],
        model_comparison=[{"model_id": "m1", "model_type": "doe_quadratic", "status": "draft",
                           "metrics": {"r2": 0.98, "rmse": 0.01, "mae": 0.008,
                                       "adj_r2": 0.97, "aic": -120.0}}],
        best_model={"model_id": "m1", "model_type": "doe_quadratic", "target": "y",
                    "equation": "y = 1.6 + 0.1*t", "status": "draft",
                    "metrics": {"r2": 0.98}, "coefficients": {"t": 0.1, "_intercept": 1.6}},
        interactions={"factors": ["a", "b"], "matrix": [[1.0, 0.2], [0.2, 1.0]]},
        sensitivity_effects={"items": [{"input": "t", "sensitivity": 0.4, "effect_size": 1.2}]},
        monte_carlo={"n_simulations": 1000, "seed": 42, "ng_count": 12, "ng_probability": 0.012,
                     "output_mean": 1.62, "output_std": 0.02, "output_median": 1.62,
                     "percentiles": {"p1": 1.56, "p5": 1.58, "p50": 1.62, "p95": 1.66, "p99": 1.68},
                     "anomaly_rankings": [{"anomaly_id": "a1", "target_input": "t",
                                           "ng_count": 8, "ng_probability": 0.008}],
                     "capability": {"pp": 1.1, "ppk": 1.0}},
        credibility={"composite": 78.0, "level": "engineering_reference",
                     "data_coverage": 0.9, "predictive_acc": 0.8,
                     "statistical_stability": 0.7, "engineering_reasonable": 0.8,
                     "validation_degree": 0.6, "extrapolation_risk": 0.2},
        recommendations=[{"type": "interaction", "priority": "high", "factors": ["a", "b"],
                          "reason": "strong interaction"}],
        process_window={"column_limits": {"t": {"min": 170.0, "max": 190.0, "center": 180.0}},
                        "basis": "model + operating range"},
        spc_results=[{"column": "y", "chart_type": "I-MR", "n_points": 90,
                      "x_mean": 1.62, "x_ucl": 1.68, "x_lcl": 1.56,
                      "mr_ucl": 0.08, "mr_mean": 0.02, "violations": 3,
                      "capability": {"cp": 1.0, "cpk": 0.9, "pp": 0.95, "ppk": 0.85,
                                     "sigma_within": 0.02, "sigma_overall": 0.021},
                      "suggestions": [{"severity": "warning", "message": "Cpk below 1.0"}]}],
        chain_trace={"steps": [{"step": "import", "entity_id": "e1", "operator": "qa",
                                "timestamp": "2026-01-01", "status": "confirmed"}]},
        gate_summary={"data_import": "confirmed"},
        approval_record={"claims": [{"claim_id": "c1", "origin_source": "stat",
                                     "evidence_status": "statistically_supported",
                                     "text": "R2 above threshold", "source_entity_ids": ["e1"]}]},
        unconfirmed_items=["model approval pending"],
        extrapolation_summary={"out_of_range_ratio": 0.02, "max_risk_score": 0.4,
                               "recommendation": "stay inside the operating range"},
        source_labels={"stat_sig": "Statistically significant"},
    )


def test_every_populated_section_has_a_sheet():
    names = sheet_names(populated_report())
    for expected in (
        "專案資訊", "Evidence", "規格", "資料品質", "欄位角色", "模型比較",
        "最終模型", "交互作用", "敏感度與效應量", "SPC", "SPC 優化建議",
        "分佈配適", "異常情境", "蒙地卡羅", "可信度", "建議製程窗口",
        "治理與追溯", "來源標籤", "實驗建議",
    ):
        assert expected in names, f"missing sheet: {expected} (have {names})"


@pytest.mark.parametrize(
    "kwargs,expected_sheet",
    [
        ({"spec": {"limits": {"lsl": 0.0, "usl": 1.0}}}, "規格"),
        ({"quality_summary": {"issue_count": 1}}, "資料品質"),
        ({"best_model": {"coefficients": {"x": 1.0}}}, "最終模型"),
        ({"distribution_fits": {"y": [{"name": "norm"}]}}, "分佈配適"),
        ({"anomalies": [{"name": "a"}]}, "異常情境"),
        ({"monte_carlo": {"n_simulations": 10}}, "蒙地卡羅"),
        ({"credibility": {"composite": 50.0}}, "可信度"),
        ({"process_window": {"column_limits": {"x": {"min": 0, "max": 1}}}}, "建議製程窗口"),
        ({"chain_trace": {"steps": [{"step": "import"}]}}, "治理與追溯"),
        ({"unconfirmed_items": ["pending"]}, "治理與追溯"),
        ({"extrapolation_summary": {"max_risk_score": 0.5}}, "治理與追溯"),
        ({"spc_results": [{"column": "y", "capability": {"cp": 1.0}}]}, "SPC"),
    ],
)
def test_section_reaches_the_workbook_on_its_own(kwargs, expected_sheet):
    """Each section must produce its sheet independently -- a section that only
    appears when some unrelated field happens to be set is still a silent loss."""
    assert expected_sheet in sheet_names(ReportData(project_name="P", **kwargs))


def test_minimal_report_still_opens():
    names = sheet_names(ReportData(project_name="Empty"))
    assert names[0] == "專案資訊"


def test_empty_input_ranges_does_not_crash_the_workbook():
    """The exact failure reported by a user: "Cannot convert {} to Excel".

    The process-definition contract carries input_ranges (column -> [low, high]) and
    other structured keys alongside 'limits'. The specification sheet wrote every key
    except 'limits' straight into a cell, and openpyxl cannot render a dict -- so an
    empty input_ranges reached the writer as {} and took the whole export down.
    """
    data = ReportData(project_name="P", spec={"limits": {"lsl": 1.0}, "input_ranges": {}})
    ExcelReportGenerator(data).generate()  # must not raise


def test_spec_input_ranges_reach_the_workbook_as_values():
    """Non-scalar spec content should be flattened into readable rows, not dropped."""
    data = ReportData(
        project_name="P",
        spec={"limits": {"lsl": 1.0, "usl": 2.0}, "input_ranges": {"t": [170.0, 190.0]}},
    )
    wb = load_workbook(BytesIO(ExcelReportGenerator(data).generate()))
    assert "規格" in wb.sheetnames
    flat = [str(c.value) for row in wb["規格"].iter_rows() for c in row if c.value is not None]
    assert any("170" in v for v in flat), flat
    assert any("190" in v for v in flat), flat


def test_nested_values_anywhere_do_not_break_the_workbook():
    """The same class of bug from every other section that writes a raw value.

    openpyxl takes scalars only, so any dict or list reaching a cell fails the whole
    export -- not just the one in the specification sheet that was reported. Each
    section below is given a structured value it could plausibly receive.
    """
    data = ReportData(
        project_name="P",
        gate_summary={"data_import": {"status": "confirmed", "detail": "ok"}},
        unconfirmed_items=["approval"],
        best_model={"model_id": "m1", "metrics": {"r2": 0.9, "extra": {"a": 1}}, "coefficients": {"x": [1.0, 2.0]}},
        extrapolation_summary={"max_risk_score": {"score": 0.4}},
        source_labels={"ai_guess": {"meaning": "AI guess"}},
    )
    ExcelReportGenerator(data).generate()  # must not raise


def test_both_renderers_present_a_structured_spec_value_the_same_way():
    """Excel and HTML must not describe the same field differently.

    Excel raised on a dict and HTML str()'d it into "{'t': [170.0, 190.0]}" -- two
    renderers, two outcomes, one report. Both now expand it into per-entry rows.
    """
    from process_intelligence_engine.reporting.html import HTMLReportGenerator

    data = ReportData(
        project_name="P",
        spec={"limits": {"lsl": 1.0, "usl": 2.0}, "input_ranges": {"t": [170.0, 190.0]}},
    )

    xlsx = load_workbook(BytesIO(ExcelReportGenerator(data).generate()))
    xls_flat = [str(c.value) for row in xlsx["規格"].iter_rows() for c in row if c.value is not None]
    assert "input_ranges.t" in xls_flat, xls_flat
    assert any("170" in v and "190" in v for v in xls_flat), xls_flat

    raw = HTMLReportGenerator(data).generate()
    page = raw.decode("utf-8") if isinstance(raw, bytes) else raw
    assert "input_ranges.t" in page, "HTML should use the same flattened label"
    assert "170" in page and "190" in page
    assert "[170.0, 190.0]" not in page, "HTML should not fall back to a Python repr"
